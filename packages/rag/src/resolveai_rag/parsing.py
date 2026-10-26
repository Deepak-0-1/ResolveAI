"""Document parsers for PDF, DOCX, TXT, and FAQ formats."""

from __future__ import annotations

import json
import re
from io import BytesIO
from pathlib import Path

from resolveai_rag.types import ParsedDocument

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".faq", ".faq.txt", ".json"}


def detect_source_type(filename: str) -> str:
    lower = filename.lower()
    if lower.endswith(".faq.txt") or lower.endswith(".faq"):
        return "faq"
    if lower.endswith(".json"):
        return "faq"
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return "pdf"
    if suffix == ".docx":
        return "docx"
    if suffix == ".txt":
        return "txt"
    raise ValueError(f"Unsupported file type: {filename}")


def parse_bytes(filename: str, data: bytes) -> ParsedDocument:
    source_type = detect_source_type(filename)
    title = Path(filename).stem
    if source_type == "pdf":
        return _parse_pdf(title, data)
    if source_type == "docx":
        return _parse_docx(title, data)
    if source_type == "faq":
        return _parse_faq(title, data, filename)
    return _parse_txt(title, data)


def _parse_pdf(title: str, data: bytes) -> ParsedDocument:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("PDF support requires pypdf. pip install pypdf") from exc

    reader = PdfReader(BytesIO(data))
    pages: list[str] = []
    for i, page in enumerate(reader.pages):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"[Page {i + 1}]\n{text}")
    body = "\n\n".join(pages).strip()
    if not body:
        raise ValueError("PDF contained no extractable text")
    return ParsedDocument(
        title=title,
        text=body,
        source_type="pdf",
        metadata={"page_count": len(reader.pages)},
    )


def _parse_docx(title: str, data: bytes) -> ParsedDocument:
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("DOCX support requires python-docx. pip install python-docx") from exc

    doc = DocxDocument(BytesIO(data))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    body = "\n\n".join(paragraphs)
    if not body:
        raise ValueError("DOCX contained no text")
    return ParsedDocument(
        title=title,
        text=body,
        source_type="docx",
        metadata={"paragraph_count": len(paragraphs)},
    )


def _parse_txt(title: str, data: bytes) -> ParsedDocument:
    body = data.decode("utf-8").strip()
    if not body:
        raise ValueError("TXT file is empty")
    return ParsedDocument(title=title, text=body, source_type="txt")


def _parse_faq(title: str, data: bytes, filename: str) -> ParsedDocument:
    raw = data.decode("utf-8").strip()
    if not raw:
        raise ValueError("FAQ file is empty")

    if filename.lower().endswith(".json"):
        payload = json.loads(raw)
        if not isinstance(payload, list):
            raise ValueError("FAQ JSON must be a list of {question, answer} objects")
        blocks: list[str] = []
        for item in payload:
            q = str(item.get("question", "")).strip()
            a = str(item.get("answer", "")).strip()
            if q and a:
                blocks.append(f"Q: {q}\nA: {a}")
        if not blocks:
            raise ValueError("FAQ JSON had no valid Q/A pairs")
        return ParsedDocument(
            title=title,
            text="\n\n".join(blocks),
            source_type="faq",
            metadata={"pair_count": len(blocks)},
        )

    # Text FAQ: Q: ... A: ... pairs
    pairs = re.findall(
        r"Q:\s*(.*?)\s*A:\s*(.*?)(?=\nQ:|\Z)",
        raw,
        flags=re.IGNORECASE | re.DOTALL,
    )
    if not pairs:
        # Fallback: treat whole file as text FAQ content
        return ParsedDocument(title=title, text=raw, source_type="faq", metadata={"pair_count": 0})

    blocks = [f"Q: {q.strip()}\nA: {a.strip()}" for q, a in pairs if q.strip() and a.strip()]
    return ParsedDocument(
        title=title,
        text="\n\n".join(blocks),
        source_type="faq",
        metadata={"pair_count": len(blocks)},
    )
